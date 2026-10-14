"""Exportación tabular de lectura: una fila por Horario y OfertaGrupo."""
import re
import unicodedata
from django.db.models import Prefetch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from academico.choices import DiaSemana
from academico.models import PlanEstudioAsignatura
from .models import Horario, OfertaGrupo

ENCABEZADOS = ["LUGAR DE DESARROLLO Y/O CENTRO TUTORIAL", "FACULTAD", "MODALIDAD",
    "PROGRAMA ACADEMICO", "SEM", "ASIGNATURA", "GRUPO", "CUPOS", "# CREDITOS",
    "PROFESOR", "DIA", "HORA INICIO", "HORA FINAL", "AULA", "SEDE", "AREA ACADEMICA"]
ANCHOS = [42, 28, 18, 35, 9, 40, 20, 10, 12, 30, 16, 14, 14, 24, 28, 28]


def valores_presentacion(fila):
    return [v.strftime("%H:%M") if hasattr(v, "strftime") else ("" if v is None else str(v)) for v in fila]


def texto_busqueda(valor):
    return "".join(c for c in unicodedata.normalize("NFD", str(valor)) if not unicodedata.combining(c)).casefold()


def filas_exportacion(datos):
    periodo = datos["periodo"]
    relaciones = OfertaGrupo.objects.filter(grupo__activo=True, grupo__periodo=periodo,
        oferta__periodo=periodo).select_related("grupo__centro_tutorial", "grupo__programa__facultad", "grupo__semestre")
    for campo in ("centro", "facultad", "programa", "semestre"):
        objeto = datos.get(campo)
        if objeto is not None:
            filtro = {"centro": "grupo__centro_tutorial", "facultad": "grupo__programa__facultad", "programa": "grupo__programa", "semestre": "grupo__semestre"}[campo]
            relaciones = relaciones.filter(**{filtro: objeto})
    horarios = Horario.objects.filter(activo=True, oferta__activo=True, oferta__periodo=periodo)
    if any(datos.get(c) is not None for c in ("centro", "facultad", "programa", "semestre")):
        horarios = horarios.filter(oferta__ofertas_grupo__in=relaciones).distinct()
    horarios = list(horarios.select_related("oferta__asignatura", "oferta__profesor", "oferta__modalidad", "sede", "aula").prefetch_related(
        Prefetch("oferta__ofertas_grupo", queryset=relaciones, to_attr="relaciones_exportacion")))
    parejas = [(h, r) for h in horarios for r in h.oferta.relaciones_exportacion]
    # El alcance completo conserva también horarios sin grupo activo, sin inventar contexto.
    if not any(datos.get(c) is not None for c in ("centro", "facultad", "programa", "semestre")):
        parejas.extend((h, None) for h in horarios if not h.oferta.relaciones_exportacion)
    programas = {r.grupo.programa_id for _, r in parejas if r is not None}
    asignaturas = {h.oferta.asignatura_id for h in horarios}
    planes = {(p.programa_id, p.semestre_id, p.asignatura_id): p for p in
        PlanEstudioAsignatura.objects.filter(activo=True, programa_id__in=programas,
            asignatura_id__in=asignaturas).select_related("area_academica")}
    dias = {codigo: i for i, (codigo, _) in enumerate(DiaSemana.choices)}
    def orden(pareja):
        h, r = pareja
        g = r.grupo if r else None
        return (g.centro_tutorial.nombre.casefold() if g else "", g.programa.facultad.nombre.casefold() if g else "",
            g.programa.nombre.casefold() if g else "", g.semestre.numero if g else 0,
            g.codigo.casefold() if g else "", dias[h.dia], h.hora_inicio, h.pk, r.pk if r else 0)
    busqueda = texto_busqueda(datos.get("busqueda", ""))
    for h, r in sorted(parejas, key=orden):
        g = r.grupo if r else None
        plan = planes.get((g.programa_id, g.semestre_id, h.oferta.asignatura_id)) if g else None
        fila = [g.centro_tutorial.nombre if g else "", g.programa.facultad.nombre if g else "",
            h.oferta.modalidad.nombre, g.programa.nombre if g else "", g.semestre.get_numero_display() if g else "",
            h.oferta.asignatura.nombre, g.codigo if g else "", r.cupos if r else "", plan.creditos if plan else "",
            str(h.oferta.profesor), h.get_dia_display(), h.hora_inicio, h.hora_fin,
            h.aula.nombre if h.aula_id else "Aula pendiente", h.sede.nombre,
            plan.area_academica.nombre if plan and plan.area_academica else ""]
        if not busqueda or busqueda in texto_busqueda("  ".join(valores_presentacion(fila))):
            yield fila


def nombre_archivo(datos):
    alcance = "Completa"
    for campo in ("centro", "facultad", "programa"):
        if datos.get(campo):
            alcance = datos[campo].nombre
    if datos.get("semestre"):
        alcance += "_Semestre_" + datos["semestre"].get_numero_display()
    def limpiar(valor):
        valor = unicodedata.normalize("NFKD", str(valor)).encode("ascii", "ignore").decode()
        return re.sub(r"[^A-Za-z0-9_-]+", "_", valor).strip("_-")[:120] or "Planificacion"
    return f"Planificacion_{limpiar(datos['periodo'].codigo)}_{limpiar(alcance)}.xlsx"


def crear_workbook(datos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Planificación"
    hoja.append(ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="172B4D")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 40
    for fila in filas_exportacion(datos):
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"  # Los datos de texto no se ejecutan como fórmulas Excel.
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        for columna in (12, 13):
            hoja.cell(hoja.max_row, columna).number_format = "hh:mm"
    for i, ancho in enumerate(ANCHOS, 1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:P{hoja.max_row}"
    return libro
