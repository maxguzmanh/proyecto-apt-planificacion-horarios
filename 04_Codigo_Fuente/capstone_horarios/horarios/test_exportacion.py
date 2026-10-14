from datetime import time
from io import BytesIO
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from academico.models import CentroTutorial, Facultad, ProgramaAcademico, ProgramaCentroTutorial, Semestre, Sede, PlanEstudioAsignatura, Modalidad
from .models import Grupo, OfertaGrupo, Horario
from .services import calcular_horas_planificadas
from .exportacion import ENCABEZADOS, filas_exportacion, nombre_archivo
from .exportacion_forms import ExportacionPlanificacionForm
from .test_validacion_docente import DatosDocencia


class ExportacionTests(DatosDocencia, TestCase):
    def setUp(self):
        self.crear_datos()
        self.centro = self.grupo.centro_tutorial
        self.programa = self.grupo.programa
        self.facultad = self.programa.facultad
        self.semestre = self.grupo.semestre
        self.semestre2 = Semestre.objects.create(numero=2)
        self.programa2 = ProgramaAcademico.objects.create(codigo="EX-P2", nombre="Programa Dos", facultad=self.facultad)
        ProgramaCentroTutorial.objects.create(programa=self.programa2, centro_tutorial=self.centro)
        self.facultad2 = Facultad.objects.create(codigo="EX-F2", nombre="Otra Facultad")
        self.centro2 = CentroTutorial.objects.create(codigo="EX-CT2", nombre="Otro Centro")
        self.programa3 = ProgramaAcademico.objects.create(codigo="EX-P3", nombre="Programa Tres", facultad=self.facultad2)
        ProgramaCentroTutorial.objects.create(programa=self.programa3, centro_tutorial=self.centro2)
        sede2 = Sede.objects.create(codigo="EX-S2", nombre="Otra Sede", centro_tutorial=self.centro2)
        self.h1 = self.crear_h(self.oferta, self.grupo, "LU", 8, pendiente=True)
        g2 = self.grupo_nuevo(self.programa, self.semestre2, self.centro, self.periodo, "EX-G2")
        self.h2 = self.crear_h(self.nueva_oferta(self.profesor), g2, "LU", 9)
        self.g3 = self.grupo_nuevo(self.programa2, self.semestre, self.centro, self.periodo, "EX-G3")
        self.h3 = self.crear_h(self.otra_oferta, self.g3, "MA", 10)
        g4 = self.grupo_nuevo(self.programa3, self.semestre, self.centro2, self.periodo, "EX-G4")
        self.h4 = self.crear_h(self.nueva_oferta(self.otro_profesor), g4, "MI", 11, pendiente=True, sede=sede2)
        self.condiciones(self.profesor, self.otro_periodo)
        g5 = self.grupo_nuevo(self.programa, self.semestre, self.centro, self.otro_periodo, "EX-G5")
        self.h5 = self.crear_h(self.nueva_oferta(self.profesor, self.otro_periodo), g5, "LU", 8)
        self.user = get_user_model().objects.create_user(username="exportador_demo")
        self.user.user_permissions.add(*Permission.objects.filter(content_type__app_label="horarios", codename__in=["view_horario", "change_horario"]))
        self.client.force_login(self.user)

    def grupo_nuevo(self, programa, semestre, centro, periodo, codigo):
        PlanEstudioAsignatura.objects.get_or_create(programa=programa, semestre=semestre, asignatura=self.asignatura, defaults={"creditos": 2})
        return Grupo.objects.create(programa=programa, semestre=semestre, centro_tutorial=centro, periodo=periodo, codigo=codigo)

    def crear_h(self, oferta, grupo, dia, hora, pendiente=False, sede=None):
        OfertaGrupo.objects.create(oferta=oferta, grupo=grupo, cupos=20)
        return Horario.objects.create(oferta=oferta, sede=sede or self.aula.sede, aula=None if pendiente else self.aula,
            dia=dia, hora_inicio=time(hora), hora_fin=time(hora + 1))

    def datos(self, nivel="periodo", **extra):
        datos = {"periodo": self.periodo.pk}
        for c in {"semestre": ("centro", "facultad", "programa", "semestre"), "programa": ("centro", "facultad", "programa"), "facultad": ("centro", "facultad"), "centro": ("centro",), "periodo": ()}[nivel]:
            datos[c] = getattr(self, c).pk
        return {**datos, **extra}

    def descargar(self, nivel="periodo", **extra):
        return self.client.get(reverse("exportar_excel"), self.datos(nivel, **extra))

    def libro(self, nivel="periodo", **extra):
        response = self.descargar(nivel, **extra)
        self.assertEqual(response.status_code, 200, response.content[:200])
        return load_workbook(BytesIO(response.content))

    def filas(self, nivel="periodo", **extra):
        libro = self.libro(nivel, **extra)
        try:
            return list(libro.active.values)[1:]
        finally:
            libro.close()

    def test_requiere_autenticacion(self):
        self.client.logout()
        for nombre in ("exportacion_planificacion", "exportar_excel", "semestres_exportacion"):
            self.assertEqual(self.client.get(reverse(nombre)).status_code, 302)

    def test_requiere_permisos_administrativos(self):
        sin_permiso = get_user_model().objects.create_user(username="sin_exportacion")
        self.client.force_login(sin_permiso)
        for nombre in ("exportacion_planificacion", "exportar_excel", "semestres_exportacion"):
            self.assertEqual(self.client.get(reverse(nombre)).status_code, 403)

    def test_profesor_con_solo_view_no_accede(self):
        docente = get_user_model().objects.create_user(username="profesor_exportacion")
        docente.user_permissions.add(Permission.objects.get(content_type__app_label="horarios", codename="view_horario"))
        self.profesor.usuario = docente; self.profesor.save()
        self.client.force_login(docente)
        self.assertEqual(self.client.get(reverse("exportacion_planificacion")).status_code, 403)
        self.assertEqual(self.descargar().status_code, 403)

    def test_periodo_obligatorio(self):
        r = self.client.get(reverse("exportar_excel"), {"nivel": "periodo"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("periodo", r.context["form"].errors)

    def test_semestre_solo_su_programa_semestre(self):
        filas = self.filas("semestre")
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][4], "I")
        self.assertEqual(filas[0][6], self.grupo.codigo)

    def test_programa_todos_sus_semestres(self):
        filas = self.filas("programa")
        self.assertEqual(len(filas), 2)
        self.assertEqual({f[4] for f in filas}, {"I", "II"})

    def test_facultad_todos_sus_programas(self):
        filas = self.filas("facultad")
        self.assertEqual(len(filas), 3)
        self.assertEqual({f[3] for f in filas}, {self.programa.nombre, self.programa2.nombre})

    def test_centro_toda_planificacion(self):
        filas = self.filas("centro")
        self.assertEqual(len(filas), 3)
        self.assertEqual({f[0] for f in filas}, {self.centro.nombre})

    def test_periodo_completo_sin_sede(self):
        filas = self.filas()
        self.assertEqual(len(filas), 4)
        self.assertEqual({f[0] for f in filas}, {self.centro.nombre, self.centro2.nombre})
        self.assertEqual({f[14] for f in filas}, {self.aula.sede.nombre, self.h4.sede.nombre})

    def test_modalidad_como_propiedad(self):
        modalidad = Modalidad.objects.create(codigo="HIB", nombre="Híbrida")
        self.oferta.modalidad = modalidad; self.oferta.save()
        self.assertIn("Híbrida", {f[2] for f in self.filas()})

    def test_sede_persistida_y_aula_pendiente(self):
        filas = self.filas()
        self.assertEqual(sum(f[13] == "Aula pendiente" for f in filas), 2)
        self.assertTrue(any(f[13] == "Aula pendiente" and f[14] == self.h4.sede.nombre for f in filas))

    def test_16_columnas_exactas(self):
        libro = self.libro()
        self.assertEqual(list(next(libro.active.values)), ENCABEZADOS)
        self.assertEqual(libro.active.max_column, 16)
        libro.close()

    def test_otro_periodo_no_aparece(self):
        self.assertEqual(len(self.filas(periodo=self.otro_periodo.pk)), 1)
        self.assertEqual(self.filas(periodo=self.otro_periodo.pk)[0][6], "EX-G5")

    def test_filename_por_nivel(self):
        for nivel in ("semestre", "programa", "facultad", "centro", "periodo"):
            with self.subTest(nivel=nivel):
                form = ExportacionPlanificacionForm(self.datos(nivel)); self.assertTrue(form.is_valid())
                filename = nombre_archivo(form.cleaned_data)
                self.assertTrue(filename.startswith("Planificacion_PER-VAL_"))
                self.assertTrue(filename.endswith(".xlsx"))
                self.assertEqual(self.descargar(nivel)["Content-Disposition"], f'attachment; filename="{filename}"')
        self.assertEqual(self.descargar()["Content-Disposition"], 'attachment; filename="Planificacion_PER-VAL_Completa.xlsx"')

    def test_filename_sanitizado(self):
        self.programa.nombre = 'Contaduría Pública / "Demo"'; self.programa.save()
        self.assertIn("Contaduria_Publica_Demo.xlsx", self.descargar("programa")["Content-Disposition"])

    def test_workbook_valido_y_formato(self):
        libro = self.libro()
        self.assertEqual(libro.sheetnames, ["Planificación"])
        hoja = libro.active
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja.auto_filter.ref, "A1:P5")
        self.assertTrue(all(c.font.bold for c in hoja[1]))
        self.assertEqual(hoja["L2"].number_format, "hh:mm")
        self.assertIsInstance(hoja["L2"].value, time)
        libro.close()

    def test_transversal_una_fila_por_grupo_sin_multiplicar_carga(self):
        antes = calcular_horas_planificadas(self.profesor, self.periodo)
        OfertaGrupo.objects.create(oferta=self.oferta, grupo=self.g3, cupos=15)
        filas = self.filas()
        self.assertEqual(len(filas), 5)
        compartidas = [f for f in filas if f[10] == "Lunes" and f[11] == time(8)]
        self.assertEqual({f[6] for f in compartidas}, {self.grupo.codigo, self.g3.codigo})
        self.assertEqual(len(self.filas("semestre")), 1)
        self.assertEqual(calcular_horas_planificadas(self.profesor, self.periodo), antes)

    def test_sin_duplicados_ni_escrituras(self):
        antes = list(Horario.objects.order_by("pk").values())
        filas = self.filas()
        self.assertEqual(len(filas), len(set(tuple(f) for f in filas)))
        self.assertEqual(list(Horario.objects.order_by("pk").values()), antes)

    def test_filtros_opcionales_acotan_contexto(self):
        self.assertEqual(len(self.filas(centro=self.centro.pk, facultad=self.facultad.pk, programa=self.programa.pk, semestre=self.semestre.pk, sede=self.aula.sede_id, modalidad=self.modalidad.pk)), 1)

    def test_contexto_incompatible_rechazado(self):
        self.assertEqual(self.descargar("programa", facultad=self.facultad2.pk).status_code, 400)
        self.assertEqual(self.descargar("semestre", semestre="999999").status_code, 400)
        self.assertEqual(self.client.get(reverse("exportar_excel"), {"periodo": self.periodo.pk, "centro": "-1"}).status_code, 400)

    def test_periodo_invalido_y_no_get(self):
        self.assertEqual(self.descargar(periodo="invalido").status_code, 400)
        self.assertEqual(self.client.post(reverse("exportar_excel"), self.datos()).status_code, 405)

    def test_pagina_sidebar_y_sin_filtros_sede_modalidad(self):
        r = self.client.get(reverse("exportacion_planificacion"))
        self.assertContains(r, ">Filtrar</button>")
        self.assertNotContains(r, 'name="nivel"')
        self.assertContains(r, 'id="sidebar-planificacion" class="sidebar-maintenance" open')
        self.assertContains(r, "Gestionar horarios")
        self.assertNotContains(r, 'name="sede"')
        self.assertNotContains(r, 'name="modalidad"')
        self.assertNotContains(self.client.get(reverse("planificacion")), "excel-export")

    def test_filtrar_muestra_tabla_sin_descargar(self):
        r = self.client.get(reverse("exportacion_planificacion"), self.datos(filtrar="1"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'id="exportacion-tabla"')
        self.assertNotIn("Content-Disposition", r)
        self.assertEqual(len(r.context["filas"]), 4)

    def test_endpoint_semestres_sin_sede(self):
        r = self.client.get(reverse("semestres_exportacion"), self.datos("programa"))
        self.assertEqual({s["id"] for s in r.json()}, {self.semestre.pk, self.semestre2.pk})
        self.assertEqual(self.client.get(reverse("semestres_exportacion"), self.datos("programa", centro=self.centro2.pk)).json(), [])

    def test_consultas_constantes(self):
        form = ExportacionPlanificacionForm(self.datos()); self.assertTrue(form.is_valid())
        with CaptureQueriesContext(connection) as antes:
            list(filas_exportacion(form.cleaned_data))
        self.crear_h(self.nueva_oferta(self.profesor), self.grupo, "JU", 12)
        with CaptureQueriesContext(connection) as despues:
            list(filas_exportacion(form.cleaned_data))
        self.assertEqual(len(antes), len(despues))

    def test_textos_no_son_formulas(self):
        self.asignatura.nombre = '=HYPERLINK("ejemplo")'; self.asignatura.save()
        libro = self.libro()
        self.assertEqual(libro.active["F2"].data_type, "s")
        libro.close()

    def test_periodo_completo_conserva_horario_sin_grupo(self):
        oferta = self.nueva_oferta(self.profesor)
        Horario.objects.create(oferta=oferta, sede=self.aula.sede, dia="VI", hora_inicio=time(12), hora_fin=time(13))
        filas = self.filas()
        self.assertEqual(len(filas), 5)
        self.assertEqual(sum(not f[6] for f in filas), 1)

    def test_orden_consistente_por_semestre_dia_hora(self):
        filas = self.filas("programa")
        self.assertEqual([f[4] for f in filas], ["I", "II"])
        self.assertEqual([f[11] for f in filas], [time(8), time(9)])

    def test_filtros_en_orden_sin_nivel(self):
        r = self.client.get(reverse("exportacion_planificacion"))
        html = r.content.decode()
        posiciones = [html.index(f'id="exportar-{c}"') for c in ("periodo", "centro", "facultad", "programa", "semestre")]
        self.assertEqual(posiciones, sorted(posiciones))
        self.assertNotContains(r, 'name="nivel"')
        self.assertNotContains(r, "Generar Excel")

    def test_preview_excel_mismo_filtro(self):
        r = self.client.get(reverse("exportacion_planificacion"), self.datos("programa", filtrar="1"))
        self.assertEqual(len(r.context["filas"]), len(self.filas("programa")))
        self.assertEqual([f["valores"] for f in r.context["filas"]], [[v.strftime("%H:%M") if isinstance(v, time) else ("" if v is None else str(v)) for v in fila] for fila in self.filas("programa")])
        self.assertIn(f"programa={self.programa.pk}", r.context["export_url"])
        self.assertNotIn("nivel=", r.context["export_url"])

    def test_busqueda_tabla_recorta_excel(self):
        filas = self.filas(busqueda=self.h4.sede.nombre)
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][14], self.h4.sede.nombre)
        self.assertEqual(self.filas(busqueda="Sin coincidencias ficticias"), [])

    def test_busqueda_larga_rechazada(self):
        self.assertEqual(self.descargar(busqueda="x" * 201).status_code, 400)

    def test_datos_html_escapados_en_resultados(self):
        self.asignatura.nombre = '<img src=x onerror="alert(1)">'; self.asignatura.save()
        r = self.client.get(reverse("exportacion_planificacion"), self.datos(filtrar="1"))
        self.assertNotContains(r, '<img src=x onerror=')
        self.assertContains(r, "&lt;img src=x")

    def test_sin_resultados_tabla_vacia_valida(self):
        for horario in Horario.objects.filter(oferta__periodo=self.periodo):
            horario.delete()
        r = self.client.get(reverse("exportacion_planificacion"), self.datos(filtrar="1"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["filas"], [])
        self.assertContains(r, 'data-row-count="0"')
        self.assertNotContains(r, 'id="exportacion-fallback"')
